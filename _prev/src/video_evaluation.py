from _tools import *
from openpyxl import load_workbook

def get_timestamp_from_sec(sec):
    second = sec%60
    minute = (sec//60)%60 
    hour = sec//60//60

    timestamp = f"{hour:02d}:{minute:02d}:{second:02d}"
    return timestamp

def get_sec_from_timestamp(timestamp:str):
    h, m, s = map(int, timestamp.split(":"))
    total_seconds = h * 3600 + m * 60 + s
    return str(total_seconds)

def convert_infer_res_xlsx_to_json(xlsx_path:str):
    wb = load_workbook(xlsx_path)
    ws = wb.active

    # 헤더 추출 (1행)
    headers = [cell.value for cell in ws[1]]

    # 데이터 추출 (2행부터)
    json_data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        row_dict["TIME"] = get_sec_from_timestamp(row_dict["TIMESTAMP"])
        row_dict["SCORE"] = round(row_dict["SCORE"], 3)
        row_dict["DESCRIPT"] = re.sub(r'[^A-Z]', '', row_dict["DESCRIPT"].upper())
        if row_dict["DESCRIPT"] == "SLEEP":
            row_dict["IS_SLEEP"] = "O"
            row_dict["DESCRIPT"] = "SLEEP"
        else:
            row_dict["IS_SLEEP"] = "X"
            row_dict["DESCRIPT"] = "AWAKE"
        # print(row_dict)
            
        json_data.append(row_dict)
    
    return json_data

if __name__ == "__main__":
    root_path = "."
    gt_root = f'{root_path}/gt/json'
    infer_root = f'{root_path}/excel'
    file_name = "류_상250426_11"
    # file_name = "조_우0427_14"
    # file_name = "baby_test_video(LGU+1)"
    # file_name = "baby_test_video(LGU+2)"
    gt_json_path = f'{gt_root}/{file_name}.json'
    infer_excel_path = f'{infer_root}/{file_name}_description_key6.xlsx'

    gt_datas = fs.load_json(gt_json_path)
    infer_results = convert_infer_res_xlsx_to_json(infer_excel_path)

    for gt_data in gt_datas:
        if gt_data["ACTIONTYPE"] != "SLEEP":
            gt_data["ACTIONTYPE"] = "AWAKE"

    eval_res = {"SLEEP" : 0, "AWAKE" : 0, "IS_SLEEP" : 0, "CHANGE" : [0, 0, 0, 0]}
    gt_exist = {"SLEEP" : 0, "AWAKE" : 0, "IS_SLEEP" : 0, "CHANGE" : [0, 0]}

    infer_idx = 0
    for gt_data in gt_datas:
        try:
            infer_result = infer_results[infer_idx]
        except:
            break
        
        if gt_data["TIMESTAMP"] != infer_result["TIMESTAMP"]:
            continue

        # print(gt_data["ACTIONTYPE"], infer_result["DESCRIPT"])

        gt_exist["IS_SLEEP"] += 1
        if gt_data["IS_SLEEP"] == infer_result["IS_SLEEP"]:
            eval_res["IS_SLEEP"] += 1

        gt_exist[gt_data["ACTIONTYPE"]] += 1
        if gt_data["ACTIONTYPE"] == infer_result["DESCRIPT"]:
            eval_res[gt_data["ACTIONTYPE"]] += 1

        infer_idx += 1
            
    # GT - Infer
    for gt_data in gt_datas:
        if gt_data["CHANGE"] == "X":
            continue

        if gt_data["ACTIONTYPE"] == "SLEEP":
            gt_exist["CHANGE"][0] += 1
        elif gt_data["ACTIONTYPE"] == "AWAKE":
            gt_exist["CHANGE"][1] += 1

        print(gt_data["TIME"])     
        correct_gap = 1
        temp_eval = [0, 0, 0, 0]   
        for inter_time in range(int(gt_data["TIME"] - correct_gap), int(gt_data["TIME"] + correct_gap) + 1):
            inter_time = str(inter_time)
            for infer_data in infer_results:
                # print(inter_time, infer_data["TIME"])
                if int(inter_time) < int(infer_data["TIME"]):
                    break
                if infer_data["TIME"] != inter_time:
                    continue
                if infer_data["CHANGE"] == "O":
                    if gt_data["ACTIONTYPE"] == "SLEEP" and infer_data["DESCRIPT"] == "SLEEP":
                        temp_eval[0] += 1
                        break
                    if gt_data["ACTIONTYPE"] == "SLEEP" and infer_data["DESCRIPT"] == "AWAKE":
                        temp_eval[1] += 1
                    if gt_data["ACTIONTYPE"] == "AWAKE" and infer_data["DESCRIPT"] == "AWAKE":
                        temp_eval[2] += 1
                        break
                    if gt_data["ACTIONTYPE"] == "AWAKE" and infer_data["DESCRIPT"] == "SLEEP":
                        temp_eval[3] += 1
            
            if temp_eval[0] > 0 or temp_eval[2] > 0:
                break
        
        if temp_eval[0] > 0:
            eval_res["CHANGE"][0] += 1
        elif temp_eval[2] > 0:
            eval_res["CHANGE"][2] += 1
        elif temp_eval[1] > 0:
            eval_res["CHANGE"][1] += 1
        elif temp_eval[3] > 0:
            eval_res["CHANGE"][3] += 1

    for key in gt_exist.keys():
        print(f'{key} : {eval_res[key]}/{gt_exist[key]}')
        
