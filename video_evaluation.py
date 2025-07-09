from _tools import *
from openpyxl import load_workbook

def get_timestamp_from_sec(sec):
    second = sec%60
    minute = (sec//60)%60 
    hour = sec//60//60

    timestamp = f"{hour:02d}:{minute:02d}:{second:02d}"
    return timestamp

def get_sec_from_timestamp(timestamp:str):
    h, m, s = map(int, timestamp.split(":"))
    total_seconds = h * 3600 + m * 60 + s
    return str(total_seconds)

def datetime_str_to_seconds(datetime_str, format):
    dt = datetime.strptime(datetime_str, format)
    return int(dt.timestamp())

def seconds_to_datetime_str(seconds, format):
    dt = datetime.fromtimestamp(seconds)
    return dt.strftime(format)

# def convert_infer_res_xlsx_to_json(xlsx_path:str):
#     wb = load_workbook(xlsx_path)
#     ws = wb.active

#     # 헤더 추출 (1행)
#     headers = [cell.value for cell in ws[1]]

#     # 데이터 추출 (2행부터)
#     json_data = []
#     for row in ws.iter_rows(min_row=2, values_only=True):
#         row_dict = dict(zip(headers, row))
#         row_dict["TIME"] = get_sec_from_timestamp(row_dict["TIMESTAMP"])
#         row_dict["SCORE"] = round(row_dict["SCORE"], 3)
#         row_dict["DESCRIPT"] = re.sub(r'[^A-Z]', '', row_dict["DESCRIPT"].upper())
#         if row_dict["DESCRIPT"] == "SLEEP":
#             row_dict["IS_SLEEP"] = "O"
#             row_dict["DESCRIPT"] = "SLEEP"
#         else:
#             row_dict["IS_SLEEP"] = "X"
#             row_dict["DESCRIPT"] = "AWAKE"
#         # print(row_dict)
            
#         json_data.append(row_dict)
    
#     return json_data

gt_to_pred_cls = {
    "baby_sleep" : "Sleep",
    "baby_awake" : "Awake",
    "baby_moving" : "Moving",
    "baby_crying" : "Crying",
    "no_baby" : "Nobaby",
    "unknown" : "Unknown",
    "baby_cough" : "Awake"
}

def gt_to_pred_json(gt_json_path):
    gt_json_datas = fs.load_json(gt_json_path)

    result = {}
    gt_video_time_diff = 0
    cnt = 0
    for idx, json_data in enumerate(gt_json_datas):
        print(json_data)
        start_time = json_data["start_time"]
        end_time = json_data["end_time"]
        event_lists = json_data["event"]
        start_sec = datetime_str_to_seconds(start_time, "%m-%d-%Y %H:%M:%S")
        end_sec = datetime_str_to_seconds(end_time, "%m-%d-%Y %H:%M:%S")

        # gt 시간이 비디오 내 적혀있는 녹화 당시 시간으로 되어있어 비디오 실행 시간으로 변환하기 위함
        if idx == 0:
            gt_video_time_diff = start_sec

        # baby_awake, baby_crying, baby_moving은 모두 baby_awake를 포함해야 하는데 그렇지 않은 케이스가 존재함
        # baby_cough는 대상이 아니므로 baby_awake로 변환함
        if event_lists[0] not in ["baby_sleep", "baby_awake", "unknown", "no_baby", "baby_cough"]:
            event_lists.insert(0, "baby_awake")
        # gt와 pred의 클래스명이 서로 상이함
        event_lists = [gt_to_pred_cls[cls] for cls in event_lists]
        
        for tidx, gt_sec in enumerate(range(start_sec, end_sec+1, 1)):
            video_sec = gt_sec - gt_video_time_diff
            # print(video_sec)
            # start time이 겹치는 gt data가 가끔 존재하여 이를 처리하기 위함
            if video_sec in result:
                continue

            temp_result = {}
            temp_result["record_start_time"] = seconds_to_datetime_str(gt_sec, "%m-%d-%Y %H:%M:%S")
            temp_result["record_end_time"] = seconds_to_datetime_str(gt_sec + 1, "%m-%d-%Y %H:%M:%S")
            temp_result["time"] = video_sec
            
            temp_result["event_full"] = event_lists
            temp_result["event"] = [event_lists[0]]

            # if tidx == 0:
            #     temp_result["scene_change"] = "O"
            # else:
            #     temp_result["scene_change"] = "X"

            result.update({video_sec : temp_result})
            # print(video_sec, result[video_sec])

    gt_dst_path = f"{root_path}/gt/json/{file_name}.json"
    fs.save_json(gt_dst_path, result)

def convert_infer_res_xlsx_to_json(xlsx_path:str):
    wb = load_workbook(xlsx_path)
    ws = wb.active

    # 헤더 추출 (1행)
    headers = [cell.value for cell in ws[1]]

    # 데이터 추출 (2행부터)
    json_data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        row_dict["TIME"] = get_sec_from_timestamp(row_dict["TIMESTAMP"])
        # row_dict["DESCRIPT"] = re.sub(r'[^A-Z]', '', row_dict["DESCRIPT"].upper())
        # if row_dict["DESCRIPT"] == "SLEEP":
        #     row_dict["IS_SLEEP"] = "O"
        #     row_dict["DESCRIPT"] = "Sleep"
        # else:
        #     row_dict["IS_SLEEP"] = "X"
        #     row_dict["DESCRIPT"] = "Awake"
        # print(row_dict)
            
        json_data.append(row_dict)
    
    return json_data

if __name__ == "__main__":
    root_path = "."
    gt_lg_root = f'{root_path}/gt/lguplus_json'
    infer_root = f'{root_path}/excel'
    file_name = "조_우0426_12"
    # file_name = "김_현0426_5"

    # gt 형식의 json 데이터를 1초 단위로 가공한 json 파일 생성
    gt_json_path = f'{gt_lg_root}/{file_name}.json'
    gt_to_pred_json(gt_json_path)

    # 비디오 필터링 평가
    gt_root = f'{root_path}/gt/json'
    gt_json_path = f'{gt_root}/{file_name}.json'
    gt_datas = fs.load_json(gt_json_path)

    # 예측 xlsx to json
    # pred_excel_path = f'{root_path}/results/excel/{file_name}.xlsx'
    # pred_datas = convert_infer_res_xlsx_to_json(pred_excel_path)
    pred_json_path = f'{root_path}/results/json/{file_name}.json'
    pred_datas = fs.load_json(pred_json_path)

    # for gt_data in gt_datas:
    #     print(gt_data)
    #     break

    # for pred_data in pred_datas:
    #     print(pred_data)
    #     break

    type_num = 4

    y_true = [data["event"] for time, data in gt_datas.items()]
    y_pred = [data["result"] for data in pred_datas]

    if type_num == 2:
        labels = ["Sleep", "Awake"]
        y_true = [data["event_full"] for time, data in gt_datas.items()]
        y_pred = [data["result"] for data in pred_datas]
        for idx, events in enumerate(y_true):
            y_temp = []
            for true in events:
                if true in ["Awake", "Moving", "Crying", "Nobaby", "Unknown"]:
                    if "Awake" in y_temp:
                        continue
                    y_temp.append("Awake")
                else:
                    y_temp.append(true)
            y_true[idx] = y_temp
        # y_true = [["Sleep"] if "Sleep" in event else ["Awake"] for event in y_true]
        y_pred = ["Sleep" if event == "Sleep" else "Awake" for event in y_pred]
    elif type_num == 4:
        labels = ["Sleep", "Awake", "Nobaby", "Unknown"]
        y_true = [data["event_full"] for time, data in gt_datas.items()]
        y_true_ori = y_true
        y_true = [["Unknown"] if "Unknown" in event else event for event in y_true]
        # y_true = [["Sleep"] if "Sleep" in event else event for event in y_true]
        # y_true = [["Awake"] if "Awake" in event else event for event in y_true]
        # y_true = [["Awake"] if "Moving" in event else event for event in y_true]
        # y_true = [["Awake"] if "Crying" in event else event for event in y_true]
        for idx, events in enumerate(y_true):
            y_temp = []
            for true in events:
                if true in ["Awake", "Moving", "Crying"]:
                    if "Awake" in y_temp:
                        continue
                    y_temp.append("Awake")
                else:
                    y_temp.append(true)
            y_true[idx] = y_temp
        # y_true = [["Sleep"] if "Sleep" in event else event for event in y_true]
        y_true = [event[0] for event in y_true]
        y_pred = [data["result"] for data in pred_datas]
        y_pred = ["Awake" if event in ["Awake", "Moving", "Crying"] else event for event in y_pred]
        print(y_true)
    elif type_num == 6:
        labels = ["Sleep", "Awake", "Moving", "Crying", "Nobaby", "Unknown"] 
        y_true = [data["event_full"] for time, data in gt_datas.items()]
        # y_true = [["Unknown"] if "Unknown" in event else event for event in y_true]
        # y_true = [["Sleep"] if "Sleep" in event else event for event in y_true]
        y_pred = [data["result"] for data in pred_datas] 

    scores = {}
    # print("GT", y_true)
    # print("Pred", y_pred)
    # print("-----------------\n")

    b_analysis = True
    text_result = []
    if b_analysis:
        cnt = 0
        for idx, (true, pred) in enumerate(zip(y_true, y_pred)):
            # print("GT", true)
            # print("Pred", pred)
            # if true != pred:
            #     print(true, pred)
            #     cnt += 1
            exclude = False
            if true not in ["Sleep", "Awake", "Unknown"]:
                exclude = True
            check = "O" if pred == true else "X"
            text_result.append(f"{y_true_ori[idx]}. {true}. {pred}. {check}. {exclude}")
            # break
        print(cnt)
        # exit()

    print(file_name)
    print("---------------------------------------------------")
    for label in labels:
        print(f"{label}(gt) : ", sum([label in true for true in y_true]))
    print("---------------------------------------------------")
    for label in labels:
        print(f"{label}(pred) : ", sum([label == pred for pred in y_pred]))
    print("---------------------------------------------------")
    
    for cls in labels:
        print("Class", cls)
        tp = sum((cls in yt and cls == yp) for yt, yp in zip(y_true, y_pred))
        fp = sum((cls not in yt and cls == yp) for yt, yp in zip(y_true, y_pred))
        fn = sum((cls in yt and cls != yp) for yt, yp in zip(y_true, y_pred))
        # fn = sum((cls in yt and cls != yp and "Unknown" != yp) for yt, yp in zip(y_true, y_pred))

        precision = round(tp / (tp + fp), 3) if (tp + fp) > 0 else 0.0
        recall = round(tp / (tp + fn), 3) if (tp + fn) > 0 else 0.0
        f1 = round((2 * precision * recall / (precision + recall)), 3) if (precision + recall) > 0 else 0.0

        scores[cls] = {
            "tp" : tp,
            "fp" : fp,
            "fn" : fn,
            "precision": precision,
            "recall": recall,
            "f1": f1
        }

        print(scores[cls])
    
    # 클래스별 support 계산
    supports = {cls: sum(1 for yt in y_true if cls in yt) for cls in labels}
    valid_classes = [cls for cls, sup in supports.items() if sup > 0]

    # # Macro
    # if valid_classes:
    #     macro_precision = sum(scores[cls]["precision"] for cls in valid_classes) / len(valid_classes)
    #     macro_recall = sum(scores[cls]["recall"] for cls in valid_classes) / len(valid_classes)
    #     macro_f1 = sum(scores[cls]["f1"] for cls in valid_classes) / len(valid_classes)
    # else:
    #     macro_precision = macro_recall = macro_f1 = 0.0

    # print("---------------------------------------------------")
    # print("[Macro]")
    # print("Macro Precision:", macro_precision)
    # print("Macro Recall:", macro_recall)
    # print("Macro F1:", macro_f1)

    # Micro
    total_tp = sum(s["tp"] for s in scores.values())
    total_fp = sum(s["fp"] for s in scores.values())
    total_fn = sum(s["fn"] for s in scores.values())

    micro_precision = total_tp / (total_tp + total_fp) if (total_tp + total_fp) > 0 else 0.0
    micro_recall = total_tp / (total_tp + total_fn) if (total_tp + total_fn) > 0 else 0.0
    micro_f1 = (2 * micro_precision * micro_recall / (micro_precision + micro_recall)) if (micro_precision + micro_recall) > 0 else 0.0

    print("---------------------------------------------------")
    print("[Micro]")
    print("Micro Precision:", round(micro_precision, 3))
    print("Micro Recall:", round(micro_recall, 3))
    print("Micro F1:", round(micro_f1, 3))

    # # Weighted
    # total_support = sum(supports[cls] for cls in valid_classes)

    # if total_support > 0:
    #     weighted_precision = sum(scores[cls]["precision"] * supports[cls] for cls in valid_classes) / total_support
    #     weighted_recall = sum(scores[cls]["recall"] * supports[cls] for cls in valid_classes) / total_support
    #     weighted_f1 = sum(scores[cls]["f1"] * supports[cls] for cls in valid_classes) / total_support
    # else:
    #     weighted_precision = weighted_recall = weighted_f1 = 0.0

    # print("---------------------------------------------------")
    # print("[Weighted]")
    # print("Weighted Precision:", weighted_precision)
    # print("Weighted Recall:", weighted_recall)
    # print("Weighted F1:", weighted_f1)
    # print("---------------------------------------------------")
        
        
